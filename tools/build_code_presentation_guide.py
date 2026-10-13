from pathlib import Path
import re
import html
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'presentation'
OUT.mkdir(parents=True, exist_ok=True)
sections = [
('ภาพรวมและคำพูดเปิดพรีเซนต์', [
('สาระหลัก', 'JUBUTSU The Seventh Room เป็นเกมสามมิติมุมมองบุคคลที่หนึ่งใน Godot ใช้ภาษา GDScript ผู้เล่นต้องจำทางเดินปกติ สังเกตความผิดปกติ และเลือกเดินต่อหรือย้อนกลับให้ถูก เพื่อไปถึงห้อง 7 ระบบใช้ฉากทางเดินเดิมซ้ำแล้วคืนค่าฉากและเปลี่ยนเหตุการณ์ในแต่ละรอบ'),
('คำพูดเปิดประมาณหนึ่งนาที', '“เกมของเราเน้นการสังเกตและจำกฎ ผู้เล่นเริ่มจากห้อง 1 ซึ่งเป็นทางเดินปกติ จากนั้นแต่ละห้องอาจมีเหตุการณ์ผิดปกติ ถ้ามีต้องย้อนกลับ ถ้าไม่มีให้เดินต่อ และห้อง 2 3 5 ต้องสั่นกระดิ่งหนึ่งครั้ง โค้ดแบ่งเป็นตัวควบคุมเกม ตัวตรวจสอบกฎ ตัวควบคุมผู้เล่น และตัวจัดการเหตุการณ์ ทำให้แก้ระบบแต่ละส่วนและทดสอบได้แยกกัน เมื่อไปถึงห้อง 7 จะได้ฉากจบตามการเก็บเครื่องราง”'),
('ไฟล์และฉากที่ควรรู้', 'project.godot กำหนดฉากเริ่มเป็น scenes/main.tscn ภายในมี SchoolLevel, Player, AnomalyDirector, HUD และ Audio ส่วน scenes/school_level.tscn เก็บฉากโรงเรียนและวัตถุ scripts เก็บพฤติกรรม data เก็บค่าความยากและรายการเหตุการณ์ assets เก็บโมเดลและเสียง tools เก็บเครื่องมือสร้างฉาก ตรวจสอบ และทดสอบ'),
('จุดที่ต้องพูดให้ตรงโค้ด', 'มีรายการเหตุการณ์ 27 รายการ แต่ ID 15 ถูกปิดด้วย enabled=false จึงเหลือ 26 รายการที่เปิดให้สุ่มในข้อมูลปัจจุบัน โอกาสมีเหตุการณ์เริ่มต้น 0.65 หรือประมาณ 65% ใช้กับห้องที่มีสิทธิ์สุ่ม ไม่ใช่ทุกห้อง ห้อง 1 ปลอดภัยและห้องสุดท้ายปกติในโหมดง่ายและปกติ'),
('วิธีใช้คู่มือนี้', 'อ่านภาพรวมก่อน แล้วฝึกอธิบายลำดับงานในหน้าถัดไป จากนั้นเลือกเปิดไฟล์และฟังก์ชันที่อ้างอิงท้ายแต่ละหน้าเวลาซ้อมตอบคำถาม หมายเลขบรรทัดอ้างอิงโค้ดในวันที่ 13 ตุลาคม 2569 และอาจเปลี่ยนเมื่อแก้ไฟล์')
], [('scripts/game.gd','_ready'), ('scripts/anomaly_director.gd','select_anomaly')]),
('ลำดับการทำงานและสถานะของเกม', [
('เริ่มระบบ', '_ready() ใน game.gd โหลดความคืบหน้า สร้างปุ่มควบคุม ส่งค่าความเร็วและ stamina ให้ Player ตั้งกระจก เรียก director.setup() และเชื่อม signal จากเมนู ผู้เล่น และเหตุการณ์ จากนั้นเตรียมห้อง 1 แต่ยังไม่เปิดให้ผู้เล่นเดิน'),
('เริ่มเล่นและอัปเดต', 'start_run() อ่านโหมด รีเซ็ตเวลา จำนวนรอบ เครื่องราง และประวัติสุ่ม แล้วเรียก load_room() และ resume_game() ระหว่างเล่น _process(delta) เพิ่มเวลา เรียก director.tick(delta) ตรวจเป้าหมายที่โต้ตอบ และตรวจว่าผู้เล่นเข้าสู่ทางออกหรือตกจากฉาก'),
('ออกจากห้อง', 'school_level.exit_at() คืน 1 เมื่อเข้าทางออกหน้า คืน -1 เมื่อเข้าทางออกหลัง และคืน 0 เมื่อยังไม่ถึงทางออก game.gd ส่ง turned_back ไป choose_exit() ถ้ายังถูกไล่ตามจะล้มเหลวก่อนตรวจ Rules.evaluate() ผล verdict มี ok, won, reason และ next เพื่อแยกการผ่านห้อง จบเกม และเริ่มใหม่'),
('ตัวแปรสถานะ', 'running บอกว่ากำลังเล่นและอนุญาตให้ระบบเหตุการณ์เดินต่อ transitioning ป้องกันอินพุตหรือการเปลี่ยนห้องซ้อน started บอกว่าเริ่มรอบเล่นแล้ว player.enabled อนุญาตการเคลื่อนไหว ส่วน hud.mode เลือกหน้าจอ เช่น title, playing, pause, note, collection และ cutscene จึงเป็นการจัดสถานะผ่านตัวแปรและข้อความ ไม่ใช่ state machine แบบ enum'),
('คำถาม ทำไมมี signal', 'ตอบว่า signal คือการแจ้งเหตุการณ์ เช่น director.danger เชื่อมกับ fail_run และ player.footstep เชื่อมกับเสียงฝีเท้า ผู้ส่งไม่ต้องรู้รายละเอียดผู้รับ อย่างไรก็ตาม director ยังอ้างถึง level, player, sound และ config โดยตรง จึงไม่ได้แยกทุกระบบออกจากกันทั้งหมด'),
('คำถาม พักเกมหยุดทุกอย่างหรือไม่', 'ตอบว่าโค้ดนี้พักโดย running=false และ player.enabled=false พร้อมพักเสียงเอฟเฟกต์ ตัวอัปเดตเกมจึงไม่เรียก tick ไม่ได้ตั้ง get_tree().paused=true เพลงยังเล่นต่อ และบางส่วนที่มี _process ของตัวเอง เช่นกระจก ยังทำงานได้')
], [('scripts/game.gd','start_run'),('scripts/game.gd','_process'),('scripts/game.gd','choose_exit'),('scripts/game.gd','pause_game')]),
('กฎผ่านห้องและกรณียกเว้น', [
('หน้าที่ของ run_rules', 'run_rules.gd เป็น RefCounted และมี static func evaluate() รับเลขห้อง ห้องสุดท้าย มีเหตุการณ์หรือไม่ ย้อนกลับหรือไม่ จำนวนกระดิ่ง ห้องที่บังคับกระดิ่ง และ ID เหตุการณ์ แล้วคืน Dictionary ผลตัดสิน ไม่ต้องเข้าถึงตำแหน่งผู้เล่นหรือฉาก จึงทดสอบกฎได้โดยป้อนค่าโดยตรง'),
('ลำดับเงื่อนไขมีผล', 'ตรวจห้องสุดท้ายก่อน โดยเดินหน้าจะชนะ จากนั้นตรวจว่าห้องบังคับกระดิ่งมี rings เท่ากับ 1 หรือไม่ ต่อด้วยข้อยกเว้น ID 27 ที่ต้องเดินหน้า แล้วจึงตรวจ anomaly != turned_back ถ้าไม่เท่ากันคือเลือกทางผิด กรณีผ่านทั่วไปคืน next=room+1'),
('ตัวอย่างสำหรับอธิบายโค้ด', 'ห้อง 2 มีเหตุการณ์ ย้อนกลับ และสั่น 1 ครั้ง จะผ่านไปห้อง 3 แต่ถ้าสั่น 2 ครั้งจะกลับห้อง 1 แม้เลือกทางถูก ห้อง 4 ปกติและเดินหน้าจะผ่าน ห้อง 4 ปกติแต่ย้อนกลับจะผิด ส่วน ID 27 ต้องเดินหน้า แม้ถือว่าเป็นเหตุการณ์ผิดปกติ และถ้าอยู่ห้อง 2 3 5 ก็ยังต้องสั่นกระดิ่ง'),
('เหตุผลที่แยกกฎจาก game', 'game.gd ดูแลอินพุต ฉากและการเปลี่ยนรอบ ส่วน evaluate() ดูแลการตัดสิน ทำให้เปลี่ยนกฎโดยไม่ต้องแก้การเคลื่อนไหว และใช้กฎเดียวกันในเกมกับชุดทดสอบ ลดปัญหากฎในเกมไม่ตรงกับกฎในเครื่องมือทดสอบ'),
('คำถาม กระดิ่งซ้ำทำไมยังนับ', 'ตอบว่า interact() เพิ่ม rings ทุกครั้งที่กดกระดิ่ง evaluate() ใช้ rings != 1 เป็นเกณฑ์ จึงทั้งไม่กดและกดซ้ำถือว่าผิดในห้องบังคับ ส่วน Audio หยุดเสียงกระดิ่งเดิมเพื่อไม่ให้เสียงซ้อน แต่ไม่ได้หยุดการนับ'),
('คำถาม ถ้าข้ามทางออกขณะถูกไล่ตาม', 'ตอบว่า choose_exit() ตรวจ chase_active ก่อน Rules.evaluate() และให้แพ้ จึงข้ามเส้นทางออกเพื่อเลี่ยงช่วงไล่ตามไม่ได้ ความอันตรายระหว่างเดินและกฎเลือกทางเป็นคนละขั้นตอน')
], [('scripts/run_rules.gd','evaluate'),('scripts/game.gd','interact'),('scripts/game.gd','choose_exit')]),
('การสุ่มและคืนค่าความผิดปกติ', [
('ข้อมูลแยกจากพฤติกรรม', 'data/anomalies.json เก็บ id, title, effect, enabled, weight, min_room และ art_status แต่พฤติกรรมจริงใน director ใช้ match id ไม่ได้เรียกพฤติกรรมจากข้อความ effect โดยอัตโนมัติ การเพิ่มแถว JSON อย่างเดียวจึงยังไม่พอ ต้องเพิ่มโค้ดและวัตถุที่เกี่ยวข้องด้วย'),
('ขั้นตอนสุ่ม', 'select_anomaly() ตรวจห้องปลอดภัยและห้องสุดท้ายก่อน ตรวจ forced_anomaly จากนั้นสุ่มว่าจะมีเหตุการณ์หรือไม่ แล้วคัดเฉพาะ enabled=true ผ่าน min_room และยังไม่อยู่ใน recent เมื่อรายชื่อที่เลือกได้หมดจะล้าง recent และสร้างรายชื่อใหม่'),
('น้ำหนักสุ่มทำงานอย่างไร', 'รวม weight ที่ไม่ติดลบเป็น total สุ่ม draw ตั้งแต่ 0 ถึง total แล้วหักน้ำหนักทีละรายการจน draw<=0 รายการนั้นถูกเลือกและเพิ่มใน recent น้ำหนัก 2 เทียบกับ 1 มีโอกาสประมาณสองเท่าเมื่ออยู่ในกลุ่มตัวเลือกเดียวกัน ปัจจุบันทุกรายการมีน้ำหนัก 1 จึงใกล้เคียงกันในกลุ่มที่ยังมีสิทธิ์'),
('ป้องกันเหตุการณ์ค้าง', 'setup() เรียก remember() เก็บ snapshot ตำแหน่ง การหมุน ขนาด visibility วัสดุ และข้อความของโหนดที่เกี่ยวข้อง reset_room() คืนค่าก่อนใช้เหตุการณ์ใหม่ พร้อมล้างเวลา chase eye_timer คืนแสง หยุดเสียง และกำหนดตัวละครใหม่ จึงไม่ปล่อยเลือดหรือวัตถุที่หายจากรอบก่อนค้าง'),
('คำถาม สุ่มซ้ำได้หรือไม่', 'ตอบว่าเหตุการณ์ที่เลือกแล้วถูกตัดออกด้วย recent จนตัวเลือกที่มีสิทธิ์หมด แล้วเริ่มชุดใหม่ ห้องปกติยังออกซ้ำได้เพราะ active_id=0 ไม่อยู่ในประวัตินี้ เมื่อแพ้กลับห้อง 1 ไม่ได้ล้าง recent แต่การเริ่มเกมใหม่ผ่าน start_run() ล้างประวัติ'),
('คำถาม อยากเพิ่มความผิดปกติใหม่', 'ตอบว่าเพิ่มข้อมูล ID และเงื่อนไขสุ่ม สร้างวัตถุใน school_level เพิ่ม match สำหรับการเริ่มเหตุการณ์หรือ tick สำหรับพฤติกรรมต่อเนื่อง ตรวจการคืนค่า และเพิ่มการทดสอบ โค้ดหลายจุดกำหนดช่วง 1 ถึง 27 ไว้ เช่นสมุดสะสมและ Debug จึงต้องขยายช่วงและข้อมูลประกอบด้วย')
], [('scripts/anomaly_director.gd','select_anomaly'),('scripts/anomaly_director.gd','remember'),('scripts/anomaly_director.gd','reset_room')]),
('ผู้เล่น กล้อง และพลังวิ่ง', [
('การเดินในสามมิติ', 'player.gd สืบทอด CharacterBody3D ใช้ Input.get_vector() อ่าน WASD แล้วคูณ transform.basis เพื่อแปลงทิศจากตัวละครเป็นโลก normalize ทิศเพื่อไม่ให้เดินทแยงเร็วกว่าเดินตรง ตั้ง velocity ตามความเร็ว เดิน 2.7 วิ่ง 5.5 หน่วยต่อวินาที และเรียก move_and_slide() ให้ระบบฟิสิกส์จัดการการชน'),
('เวลาและแรงโน้มถ่วง', '_physics_process(delta) ดูแลการเคลื่อนไหวและแรงโน้มถ่วงซึ่งลด velocity.y ด้วย 18*delta ส่วนการอัปเดตเหตุการณ์ทั่วไปอยู่ _process(delta) ค่า delta คือเวลาที่ผ่านตั้งแต่การอัปเดตครั้งก่อน ใช้กับเวลา แรงเร่ง และการขยับเอง ส่วน move_and_slide() รับ velocity ต่อวินาที จึงไม่คูณ velocity ด้วย delta ซ้ำ'),
('กล้องและการขึ้นพื้นต่างระดับ', 'เมาส์หมุนตัวละครในแกน Y และหมุนกล้องในแกน X โดย clamp จำกัดการเงยก้ม try_step_up() ตรวจเส้นทางชนและยิง ray ไปพื้นเพื่อขึ้นขั้นเตี้ยได้ไม่เกิน 0.26 หน่วย update_eye_position() ตรวจผนังหน้าใบหน้าและลดระยะยื่นกล้อง ป้องกันการมองทะลุกำแพง'),
('stamina ในโหมดปกติ', 'ความจุเริ่มต้น 6 หน่วย วิ่งลดด้วย delta จึงวิ่งต่อเนื่องได้ประมาณ 6 วินาที เมื่อหยุดรอ 1.5 วินาทีก่อนฟื้นในอัตรา 0.8 หน่วยต่อวินาที ถ้าหมดจะตั้ง exhausted และต้องฟื้นถึง 25% จึงวิ่งได้อีก คำนวณ recovery_time เฉพาะส่วนเวลาที่เลยช่วงรอ โหมดง่ายและไร้สิ้นสุดปิดการจำกัด stamina เปลี่ยนห้อง reset_stamina()'),
('คำถาม ทำไมวิ่งชนผนังยังเสียพลัง', 'ตอบว่า wants_sprint ตรวจปุ่ม Shift และอินพุตเดิน ไม่ได้ใช้ความเร็วเคลื่อนที่จริงในการลด stamina จึงอาจเสียพลังแม้ชนผนัง ส่วนแอนิเมชันและเสียงฝีเท้าใช้ get_real_velocity() เพื่อไม่ให้ภาพเดินต่อขณะยืนติดผนัง'),
('คำถาม ภาพตัวเราแยกจากกระจกอย่างไร', 'ตอบว่าตัวเต็มอยู่ layer ค่า 8 และ mesh มุมมองบุคคลที่หนึ่งอยู่ค่า 16 กล้องผู้เล่นเลือก 1|2|16 ส่วนกล้องกระจกเลือก 1|8 ในสภาพปกติ เป็น bitmask กำหนดสิ่งที่แต่ละกล้องเห็น ค่า 8 และ 16 ไม่ได้หมายถึง layer ลำดับที่ 8 และ 16')
], [('scripts/player.gd','_physics_process'),('scripts/player.gd','update_stamina'),('scripts/player.gd','try_step_up'),('scripts/player.gd','_ready')]),
('การโต้ตอบและพฤติกรรมอันตราย', [
('หาเป้าหมายกด E', 'update_interaction() ตรวจ Bell, Note, Walker/Offer, Door/Handle และ CharmPickup โดยวัตถุต้องมองเห็น อยู่ใกล้กว่า 3 หน่วย อยู่ในมุมมองที่ dot product มากกว่า 0.87 และไม่มีผนังบัง ray จึงเลือกวัตถุที่ใกล้ที่สุด ไม่ใช่การกด E แล้วรับทุกวัตถุรอบตัว'),
('สิ่งที่เกิดเมื่อโต้ตอบ', 'interact() ใช้ match target เครื่องรางตั้ง charm_collected=true และซ่อนวัตถุ กระดิ่งเพิ่ม rings และเล่นเสียง กระดาษเปิดหน้าอ่านกฎ มือจับเปิดหรือปิดประตูยกเว้น ID 3 ส่วนรับของจาก NPC เรียก fail_run() ทันที เครื่องรางไม่แสดง prompt [E] แต่ยังเก็บได้ถ้าผ่านการเลือกเป้าหมาย'),
('การสบตา', 'check_eyes() ตรวจว่า NPC มองมาทางผู้เล่นด้วย dot>0.45 และกล้องมองใกล้หัวด้วย dot>0.993 ในระยะต่ำกว่า 4 หน่วย พร้อม ray ตรวจสิ่งบัง แล้วสะสม eye_timer จนครบ 1.8 วินาที เมื่อเลิกมองเวลาลดด้วย delta*2 ไม่ได้รีเซ็ตทันที ID 8 ยกเว้นเงื่อนไขสบตา และ ID 27 ไม่เรียก check_eyes()'),
('การไล่ตาม ID 11 และ 21', 'เมื่อผ่านตำแหน่ง trigger จะสร้างช่วงไล่ตาม pursuer ใช้ move_toward() ตามตำแหน่งผู้เล่นด้วยความเร็ว 3.7 ถ้าระยะต่ำกว่า 0.48 จะส่ง danger ถ้าอยู่รอดครบ 3 วินาที pursuer หายและ chase_finished=true ไม่ได้ตรวจว่าต้องกด Shift จึงตัดสินจากการอยู่รอดและระยะจริง ผู้ไล่ตามนี้ไม่ได้ใช้ NavigationAgent3D หาเส้นทาง'),
('ข้อยกเว้น ID 27', 'ปิด sprint_allowed ตั้งแต่เริ่มห้อง เมื่อผ่าน trigger จะมีคนเดินตามหลังด้วย 90% ของความเร็วเดิน หากย้อนจากตำแหน่งที่คืบหน้าเกิน 1 หน่วยหรือถูกตามใกล้กว่า 0.6 จะแพ้ สุดท้ายต้องออกด้านหน้า ใช้กฎต่างจากเหตุการณ์อื่น'),
('คำถาม มีข้อความเตือนจาก signal หรือไม่', 'ตอบว่า director ประกาศ warning และส่งข้อความบางช่วง แต่ game._ready() ปัจจุบันเชื่อม danger เท่านั้น ไม่ได้เชื่อม warning กับ HUD จึงไม่ควรอ้างว่าผู้เล่นเห็นข้อความเตือนเหล่านี้ทุกครั้ง')
], [('scripts/game.gd','update_interaction'),('scripts/game.gd','interact'),('scripts/anomaly_director.gd','tick'),('scripts/anomaly_director.gd','check_eyes'),('scripts/anomaly_director.gd','tick_no_running')]),
('ฉากโรงเรียน กระจก และแอนิเมชัน', [
('ขอบเขตของฉาก', 'school_level.gd จัดวัสดุโมเดล แสง เลขห้อง ประตู จุดออกและกระจก prop(path) ใช้เข้าถึงวัตถุใต้ Props ผ่านชื่อที่กำหนด exit_at() แปลงตำแหน่งผู้เล่นเป็นพิกัดของทางออกและตรวจขนาดทั้ง X Y Z จึงไม่ผ่านรอบเพียงเพราะอยู่ค่า Z เดียวกันบนอีกชั้น'),
('กระจกทำงานจริง', 'setup_mirror() สร้าง SubViewport ขนาด 384x780 และ Camera3D อีกตัวซึ่งใช้โลกสามมิติเดียวกัน เอา texture ของ viewport ไปติดผิวกระจก ภาพจึงเปลี่ยนตามตำแหน่งผู้เล่น ไม่ใช่ภาพนิ่ง _process() คำนวณ reflected_eye.x=2*center.x-eye.x และใช้ set_frustum() ปรับมุมมองตามตำแหน่งตา'),
('ภาพผีและประสิทธิภาพ', 'set_ghost_reflection() เปลี่ยน cull_mask และแสดง reflection_avatar ที่ตามตำแหน่งผู้เล่นเพื่อแทนร่างในกระจก เมื่อระยะตามแกน Z จากกระจกเกิน 8 หน่วย จะปิดการอัปเดต viewport ลดงานเรนเดอร์ วิธีสะท้อนนี้อ้างอิงกระจกแนวที่เขียนไว้ในฉาก จึงไม่ใช่ระบบกระจกหมุนอิสระทุกแนว'),
('ตัวละครและท่าทาง', 'actor_visual.gd เลือก PackedScene ของ NPC ตาม variant เช่น normal, masked, ghost และ attacker เก็บโหนด Head และ Offer เป็นจุดอ้างอิงเกม แม้เปลี่ยนโมเดลจริง หน้ากากใช้ BoneAttachment3D ติดกระดูกศีรษะ animate() ปรับอัตราคลิปให้สัมพันธ์กับความเร็วและเรียก advance() ด้วยตัวเอง'),
('ท่าผู้เล่น', 'player_motion.gd สร้าง Idle Walk Run สำหรับ rig 9 กระดูกด้วยค่า sin ของรอบการเคลื่อนไหวและบันทึก Quaternion ลง animation track ส่วน player.update_body() เลือกท่าตามความเร็ว ปรับ transition และเล่นย้อนคลิปเมื่อเดินถอยหลัง จึงควรเรียกว่าท่าพื้นฐานที่สร้างด้วยโค้ด ไม่ใช่ motion capture'),
('คำถาม ทำไมวัตถุไม่จมผนัง', 'ตอบว่า prop_mounting.gd รวม AABB ของ mesh ที่แปลงตำแหน่งแล้ว ใช้ขอบจริงของโมเดลจัดให้ชิดผนังและเว้น CLEARANCE เล็กน้อย การเปลี่ยนโมเดลจึงไม่ต้องอาศัยการเดาจุดกึ่งกลางอย่างเดียว')
], [('scripts/school_level.gd','exit_at'),('scripts/school_level.gd','setup_mirror'),('scripts/school_level.gd','_process'),('scripts/actor_visual.gd','animate'),('scripts/player_motion.gd','configure'),('scripts/prop_mounting.gd','mount_right')]),
('เสียง เมนู ภาษา และข้อมูลบันทึก', [
('เสียงมีสองลักษณะ', 'audio_manager.gd ใช้ AudioStreamPlayer สำหรับเพลงและเสียงหน้าจอ ใช้ AudioStreamPlayer3D สำหรับเสียงในตำแหน่งจริง เช่นกระดิ่ง ฝีเท้าและเคาะประตู play_at() กำหนดตำแหน่ง volume pitch และ max_distance เมื่อเสียงจบใช้ finished เพื่อลบตัวเล่นเสียง เปลี่ยนห้องเรียก stop_events() เพื่อล้างเสียงเหตุการณ์'),
('การตั้งค่าและ HUD', 'hud.gd สืบทอด CanvasLayer และสร้างหน้าตาด้วยโค้ด เช่น Label Button Container และ ProgressBar ปุ่มส่ง signal กลับ game แยกเสียงเพลงผ่าน bus Music เพื่อปรับได้แยกจากระดับรวม การตั้งค่าเสียง เมาส์ ความละเอียด เต็มจอ ความสว่าง และภาษาอยู่ user://settings.cfg'),
('ภาษาและความสว่าง', 'localization.gd มี static language และ Dictionary สำหรับแปล t(source) ถ้าเลือกอังกฤษจะหาคำแปล ถ้าไม่พบใช้ข้อความต้นฉบับ rebuild_language() สร้าง UI ใหม่และรักษาโหมดที่เลือก ส่วนความสว่างใช้ shader อ่าน screen_texture แล้วปรับ pow ของสีบนหน้าจอ ไม่ใช่เปลี่ยนพลังไฟทุกดวงในฉาก'),
('ข้อมูลความคืบหน้า', 'mode_progress.gd โหลดและบันทึก ConfigFile ที่ user://progress.cfg เก็บการปลดล็อก Endless สถิติ จำนวนชนิดที่เคยพบ และฉากจบที่ปลดล็อก ตอนโหลดกรอง ID ให้ถูกช่วงและไม่ซ้ำ discover() กับ collect_ending() บันทึกเมื่อมีรายการใหม่ ต่างจากสถานะเครื่องรางที่อยู่ในหน่วยความจำของรอบเล่น'),
('คำถาม res และ user ต่างกันอย่างไร', 'ตอบว่า res:// อ้างอิงทรัพยากรในโปรเจกต์ ส่วน user:// เป็นพื้นที่เก็บข้อมูลของแอปสำหรับผู้เล่น การแยกนี้ทำให้การตั้งค่าและความคืบหน้าบันทึกข้ามการเปิดเกมได้โดยไม่ต้องแก้ไฟล์ทรัพยากรที่ส่งออกกับเกม'),
('คำถาม สมุดสะสมถือว่าพบเมื่อไร', 'ตอบว่า collect_encounter() เรียกเมื่อออกจากห้อง แพ้ หรือกลับเมนูในเงื่อนไขที่กำหนด หาก ranked_run และ started จะบันทึก active_id แม้เลือกทางผิด ไม่ได้ใช้การตรวจว่าผู้เล่นมองเห็นสิ่งผิดปกติจริง และยังไม่มีระบบเซฟตำแหน่งหรือเล่นต่อกลางรอบ')
], [('scripts/audio_manager.gd','play_at'),('scripts/hud.gd','save_preferences'),('scripts/hud.gd','build_brightness_filter'),('scripts/localization.gd','t'),('scripts/mode_progress.gd','load_progress'),('scripts/game.gd','collect_encounter')]),
('โหมดเกม เครื่องราง และฉากจบ', [
('สามโหมดใช้ระบบหลักร่วมกัน', 'ง่ายมีข้อความชวนสังเกตและไม่จำกัด stamina ปกติใช้ stamina ไร้สิ้นสุดต้องปลดล็อกหลังจบเกมที่นับผลจริง และแพ้แล้วจบรอบแทนกลับห้อง 1 start_run() ตั้ง director.endless_mode และ player.stamina_enabled ตามโหมด'),
('กฎโหมดไร้สิ้นสุด', 'choose_exit() แปลงห้องจริงเป็น rule_room=((room-1)%final_room)+1 เพื่อวนกฎกระดิ่ง และส่ง rule_final=final_room+1 ทำให้กฎห้องสุดท้ายไม่ตัดจบ รอบจริงยังเพิ่มไปเรื่อย ๆ การสุ่มไม่บังคับห้อง 7 ให้ปกติในโหมดนี้ แต่ห้องแรกยังปลอดภัยตาม config'),
('เครื่องรางและผลจบ', 'reset_charm() สุ่มห้อง 1 ถึง final_room และตำแหน่งบนทางเดิน load_room() แสดงเครื่องรางเฉพาะห้องนั้นและยังไม่เก็บ เมื่อชนะ win_run() เลือก ending="true" ถ้า charm_collected มิฉะนั้น "normal" หากแพ้หรือเริ่มรอบใหม่เครื่องรางถูกล้างและสุ่มใหม่ โหมดไร้สิ้นสุดไม่แสดงเครื่องราง'),
('คัทซีนและดูซ้ำ', 'ending_cutscene.gd รับชนิดฉากจบจาก begin() แล้วจัดตัวละคร กล้อง แสง เสียงและคำบรรยายตามเวลา ส่ง finished เมื่อจบ การ replay_ending() สร้าง school_level อีกชุดและสำเนาทรัพยากรที่ต้องเปลี่ยน แล้วคืนฉากเดิมหลังจบ เพื่อไม่เปลี่ยนห้องและของที่มีอยู่ในรอบเดิม'),
('คำถาม ทำอย่างไรไม่ให้โหมดทดสอบปลดล็อกของ', 'ตอบว่า ranked_run เริ่มจาก forced_anomaly<0 และ debug_load() ตั้งเป็น false การบันทึกฉากจบ สมุดสะสมและสถิติมีเงื่อนไขนี้กำกับ ทำให้รอบที่บังคับเหตุการณ์หรือใช้ Event Lab ไม่ถูกนับเป็นผลการเล่นปกติ'),
('คำถาม สถิติเลือกคะแนนอย่างไร', 'ตอบว่า record() เลือกรอบที่ผ่านห้องได้มากกว่า หากจำนวนห้องเท่ากันและมากกว่า 0 จะเลือกรอบที่ใช้เวลาน้อยกว่า แล้วเก็บจำนวนชนิดเหตุการณ์ของรอบนั้นด้วย ไม่ได้จัดอันดับโดยจำนวนเหตุการณ์เป็นหลัก')
], [('scripts/game.gd','reset_charm'),('scripts/game.gd','choose_exit'),('scripts/game.gd','win_run'),('scripts/game.gd','replay_ending'),('scripts/mode_progress.gd','record'),('scripts/ending_cutscene.gd','begin')]),
('การทดสอบ ข้อจำกัด และแนวตอบอาจารย์', [
('หลักฐานที่เปิดอธิบายได้', 'tools/test_game.gd ทดสอบคู่ค่าห้อง 1 ถึง 6 มีหรือไม่มีเหตุการณ์ เดินหน้าหรือย้อน และกระดิ่ง 0 1 2 ครั้ง รวม 72 กรณีพื้นฐานก่อนตรวจส่วนอื่น ยังทดสอบห้องสุดท้าย การคืนค่าฉาก ปุ่มกระดิ่ง การเคลื่อนที่ การชนและการพัก มีชุดแยก test_stamina, test_modes, test_collection, test_settings, test_seven_endings และ test_player_motion'),
('ควรอธิบายวิธีตรวจสอบ', 'แยกการทดสอบกฎแบบป้อนค่ากับการทดสอบที่สร้าง main.tscn จริง กรณีแรกช่วยตรวจตรรกะได้เร็ว กรณีหลังตรวจการทำงานร่วมกับฉาก แต่ยังควรเล่นและตรวจภาพเสียงจริง เพราะ headless test ไม่ยืนยันความสวยงามหรือความรู้สึกของเกม คู่มือนี้อธิบายชุดทดสอบจากโค้ด ไม่ใช่รายงานว่ารันทดสอบใหม่ทั้งหมดแล้ว'),
('คำถาม จุดเด่นทางเทคนิคคืออะไร', 'ตอบว่าเราแยกกฎเป็นฟังก์ชันที่ทดสอบได้ ใช้ข้อมูลกำหนดสิทธิ์และน้ำหนักสุ่ม ใช้ snapshot คืนค่าฉากก่อนแต่ละรอบ และใช้ signal แจ้งเหตุการณ์ระหว่างระบบ ยกตัวอย่าง evaluate() กับ reset_room() ให้เห็นจริงแทนการพูดชื่อแนวคิดอย่างเดียว'),
('คำถาม ข้อจำกัดและสิ่งที่จะพัฒนาต่อ', 'ตอบว่าเป็นต้นแบบที่ผูกกับชื่อโหนดและตำแหน่งฉาก การเพิ่มเหตุการณ์ยังต้องแก้ match id ผู้ไล่ตามเคลื่อนตรงโดยไม่มีระบบหาเส้นทาง บางภาพและท่าทางยังเป็นต้นแบบ ยังไม่มีเซฟกลางรอบและยังควรทดสอบบนเครื่องเป้าหมาย หากพัฒนาต่อจะตรวจความถูกต้องของ JSON และแยกพฤติกรรมเหตุการณ์ให้ขยายง่ายขึ้น'),
('ลำดับเดโมแนะนำ', 'เริ่มจากเมนูและห้อง 1 อธิบายการจำฉาก เปิดกระดาษกฎ แสดงการเดิน กล้องและกด E แล้วเปิด run_rules.evaluate() อธิบายหนึ่งตัวอย่าง จากนั้นเปิด select_anomaly() และ reset_room() หากใช้ F3 ต้องเปิด debug_enabled ใน Inspector ก่อน และแจ้งว่าเป็นรอบทดสอบที่ไม่นับปลดล็อก สุดท้ายแสดงสมุดสะสมหรือฉากจบที่ปลดล็อกแล้ว'),
('รายการทวนก่อนขึ้นพูด', 'จำไฟล์ game, run_rules, player และ anomaly_director ให้ได้ อธิบายเส้นทาง start_run → load_room → tick → choose_exit → evaluate → transition หรือ win ยกตัวอย่างกฎที่ผ่านและผิดอย่างละหนึ่งกรณี จำข้อยกเว้น ID 27 และการปิด ID 15 เตรียมตอบเรื่อง delta, signal, ray, snapshot และ user:// พร้อมเปิดไฟล์ที่เกี่ยวข้อง')
], [('tools/test_game.gd','run'),('scripts/game.gd','debug_load')]),
]

def reference(file, func):
    lines=(ROOT/file).read_text(encoding='utf-8-sig').splitlines()
    pat=re.compile(r'^(?:static )?func '+re.escape(func)+r'\(')
    n=next((i for i,line in enumerate(lines,1) if pat.match(line)),None)
    if n is None:
        raise ValueError((file,func))
    return f'{file}:{n}  {func}()'

doc=Document()
sec=doc.sections[0]
sec.page_height=Cm(29.7); sec.page_width=Cm(21)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Tahoma'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Tahoma')
    style._element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Tahoma')
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(10.5 if name=='Normal' else (20 if name=='Title' else 15 if name=='Heading 1' else 11))
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.widow_control=True
md=['# คู่มืออธิบายโค้ด JUBUTSU สำหรับเตรียมพรีเซนต์','', 'อ้างอิงโค้ดวันที่ 13 ตุลาคม 2569','']
for i,(title,parts,refs) in enumerate(sections,1):
    if i>1: doc.add_page_break()
    if i==1:
        doc.add_paragraph('คู่มืออธิบายโค้ด JUBUTSU', 'Title')
        doc.add_paragraph('สำหรับเตรียมพรีเซนต์และตอบคำถามอาจารย์\nอ้างอิงโค้ดวันที่ 13 ตุลาคม 2569')
    doc.add_paragraph(f'{i} {title}', 'Heading 1')
    md.extend([f'## {i} {title}',''])
    for label,body in parts:
        doc.add_paragraph(label,'Heading 2')
        doc.add_paragraph(body)
        md.extend([f'### {label}',body,''])
    p=doc.add_paragraph()
    r=p.add_run('เปิดโค้ดประกอบ\n');r.bold=True
    for file,func in refs:
        p.add_run(reference(file,func)+'\n').font.size=Pt(8)
    md.extend(['### เปิดโค้ดประกอบ']+[f'- `{reference(f,fn)}`' for f,fn in refs]+[''])
footer=sec.footer.paragraphs[0]
footer.alignment=2
footer.add_run('JUBUTSU  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.save(OUT/'JUBUTSU_Code_Presentation_TH.docx')
(OUT/'JUBUTSU_Code_Presentation_TH.md').write_text('\n'.join(md),encoding='utf-8')
pages=[]
for i,(title,parts,refs) in enumerate(sections,1):
    intro='<h1>คู่มืออธิบายโค้ด JUBUTSU</h1><p>สำหรับเตรียมพรีเซนต์และตอบคำถามอาจารย์<br>อ้างอิงโค้ดวันที่ 13 ตุลาคม 2569</p>' if i==1 else ''
    blocks=''.join('<h3>'+html.escape(label)+'</h3><p>'+html.escape(body)+'</p>' for label,body in parts)
    sources='<h3>เปิดโค้ดประกอบ</h3><ul>'+''.join('<li><code>'+html.escape(reference(f,fn))+'</code></li>' for f,fn in refs)+'</ul>'
    pages.append('<section>'+intro+f'<h2>{i} '+html.escape(title)+'</h2>'+blocks+sources+'</section>')
page='''<!doctype html><html lang="th"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>คู่มืออธิบายโค้ด JUBUTSU</title><style>
body{font-family:Tahoma,sans-serif;background:#eee;color:#111;margin:0;line-height:1.75}main{max-width:850px;margin:24px auto}section{background:white;padding:40px 52px;margin-bottom:24px;border-radius:8px}h1{font-size:28px}h2{font-size:22px}h3{font-size:16px;margin:20px 0 4px}p{margin:0 0 12px}code{font-size:12px;overflow-wrap:anywhere}nav{background:#fff;padding:20px 32px;margin-bottom:24px}a{color:#185391} @media(max-width:650px){section{padding:24px}main{margin:12px}h1{font-size:24px}} @media print{body{background:white;font-size:10pt;line-height:1.5}main{max-width:none;margin:0}nav{display:none}section{padding:0;border-radius:0;break-before:page}section:first-of-type{break-before:auto}h3{break-after:avoid}p{orphans:3;widows:3}@page{size:A4;margin:18mm}}
</style><main><nav>คู่มือเตรียมตอบคำถาม 10 หัวข้อ • เปิดอ่านหรือใช้คำสั่งพิมพ์ของเบราว์เซอร์</nav>'''+''.join(pages)+'</main></html>'
(OUT/'JUBUTSU_Code_Presentation_TH.html').write_text(page,encoding='utf-8')
print('Created guide:',len(sections),'sections with verified source references')
